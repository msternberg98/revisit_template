"""Minima-/Maxima-Aufgaben: absolute und normalisierte Fehler, Rangvergleiche, Diagnostik.
Datei neben Nutzerstudie_all_tidy.csv ablegen; Ausgabe: Analysis_Output/Identifikation.
Benötigt: pandas, numpy, matplotlib, openpyxl, scipy.
"""
from pathlib import Path
import re
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy import stats
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
# 1. Konfiguration
CSV_FILENAME = "Nutzerstudie_all_tidy.csv"  # Zum Testen: "Nutzerstudie_all_tidy_latin_square.csv"
# CSV_FILENAME = "Nutzerstudie_all_tidy_Test.csv"
INCLUDE_INCOMPLETE = False  # Wird nicht verwendet: nur completed, niemals in progress
OUTPUT_FOLDER = Path("Analysis_Output") / "Identifikation"
SHOW_PLOTS = False
TOLERANCE_PERCENT = 1.0  # Nur deskriptiv: Fehler <= 1 % der Variablenspanne.
METHODS = ["VSUP", "ScaledGlyph", "IsoGlyph"]
METHOD_LABELS = {"VSUP": "VSUP", "ScaledGlyph": "Scaled Glyph", "IsoGlyph": "IsoGlyph"}
COLORS = {"VSUP": "#3676b6", "ScaledGlyph": "#e5a343", "IsoGlyph": "#779b7b"}
TASKS = ["Niederschlag_Uncertainty_Maxima", "Niederschlag_Uncertainty_Minima",
         "Luftdruck_Value_Maxima", "Luftdruck_Value_Minima"]
TASK_LABELS = {"Niederschlag_Uncertainty_Maxima": "Unsicherheit: Maximum",
               "Niederschlag_Uncertainty_Minima": "Unsicherheit: Minimum",
               "Luftdruck_Value_Maxima": "Mittelwert: Maximum",
               "Luftdruck_Value_Minima": "Mittelwert: Minimum"}
TRIAL_PATTERN = re.compile(r"^(Niederschlag_Uncertainty_(?:Maxima|Minima)|Luftdruck_Value_(?:Maxima|Minima))_(VSUP|ScaledGlyph|IsoGlyph)$")
def save_plot(fig, folder, name, legend=False):
    """PNG und PDF speichern; Legende gegebenenfalls fest rechts außerhalb."""
    fig.tight_layout(rect=(0, 0, .82, 1) if legend else (0, 0, 1, 1))
    fig.savefig(folder / f"{name}.png", dpi=300, bbox_inches="tight")
    fig.savefig(folder / f"{name}.pdf", bbox_inches="tight")
    if SHOW_PLOTS:
        plt.show()
    plt.close(fig)
def describe(values):
    """Deskriptive Kennzahlen ohne automatische Entfernung statistischer Ausreißer."""
    s = pd.Series(values).dropna().astype(float)
    if s.empty:
        return {"n": 0, "Mittelwert": np.nan, "Median": np.nan, "Standardabweichung": np.nan,
                "Q1": np.nan, "Q3": np.nan, "IQR": np.nan}
    q1, q3 = s.quantile([.25, .75])
    return {"n": len(s), "Mittelwert": s.mean(), "Median": s.median(),
            "Standardabweichung": s.std(ddof=1) if len(s) > 1 else np.nan,
            "Q1": q1, "Q3": q3, "IQR": q3-q1}
def parse_trials(df):
    """Numerische Klickantworten plus zugeordnete separate Sicherheits-Likerts."""
    # Likert-Antworten liegen als separate responseId-Zeilen derselben Trial-ID vor.
    likert_rows = df[df['responseId'].astype(str).str.endswith('_Likert_Response', na=False)].copy()
    likert_rows = likert_rows[likert_rows['trialId'].astype(str).str.fullmatch(TRIAL_PATTERN.pattern, na=False)]
    likert_rows['Likert_Response'] = pd.to_numeric(likert_rows['answer'], errors='coerce')
    # Zulässiger Wertebereich der Sicherheitsfrage: 1 bis 7.
    likert_rows.loc[~likert_rows['Likert_Response'].between(1, 7) |
                    (likert_rows['Likert_Response'] % 1 != 0), 'Likert_Response'] = np.nan
    likert_rows = likert_rows.dropna(subset=['Likert_Response'])
    if likert_rows.duplicated(['participantId', 'trialId']).any():
        raise ValueError('Mehrere Likert-Antworten je Person und Aufgabe/Methode gefunden.')
    likert_lookup = likert_rows.set_index(['participantId', 'trialId'])['Likert_Response'].to_dict()
    rows = []

    for r in df.to_dict("records"):
        match = TRIAL_PATTERN.fullmatch (str(r.get ("trialId", "")))
        if not match:
            continue
        task, method = match.groups()
        correct = pd.to_numeric(r.get("correctAnswer"), errors="coerce")
        if pd.isna(correct):
            continue  # Die weiteren Trial-Zeilen sind keine Identifikationsantworten.
        raw = pd.to_numeric(r.get("answer"), errors="coerce")
        valid = pd.notna(raw) and np.isfinite(raw) and raw != -999
        duration = pd.to_numeric(r.get("cleanedDuration"), errors="coerce")
        if pd.isna(duration):
            duration = pd.to_numeric(r.get("duration"), errors="coerce")
        rows.append({"Teilnehmer": str(r["participantId"]), "Aufgabe": task,
                     "Methode": method, "Antwort": float(raw) if valid else np.nan,
                     "Sollwert": float(correct), "Likert_Response": likert_lookup.get((r["participantId"], r["trialId"]), np.nan), "Gültig": bool(valid),
                     "Ungültigkeitsgrund": "" if valid else ("-999 (keine Antwort)" if raw == -999 else "Fehlend/nicht numerisch"),
                     "Zeit_s": float(duration)/1000 if pd.notna(duration) and duration >= 0 else np.nan,
                     "trialOrder": r.get("trialOrder")})
    out = pd.DataFrame(rows)
    if out.empty:
        raise ValueError("Keine Identifikationsaufgaben gefunden. trialId und correctAnswer prüfen.")
    duplicates = out.duplicated(["Teilnehmer", "Aufgabe", "Methode"], keep=False)
    if duplicates.any():
        raise ValueError("Doppelte Identifikationsantworten gefunden; Rohdaten prüfen:\n" +
                         out.loc[duplicates, ["Teilnehmer", "Aufgabe", "Methode"]].to_string(index=False))
    # Referenzspanne stammt aus den beiden tatsächlichen Extrema der jeweiligen Variable.
    out["Variable"] = out["Aufgabe"].str.replace(r"_(?:Maxima|Minima)$", "", regex=True)
    reference = out.groupby(["Variable", "Aufgabe"])["Sollwert"].agg(["min", "max"])
    if (reference["max"]-reference["min"] > 1e-9).any():
        raise ValueError("Widersprüchliche Sollwerte für dieselbe Aufgabe: Rohdaten prüfen.")
    spans = out.groupby("Variable")["Sollwert"].agg(["min", "max"])
    spans["Spanne"] = spans["max"]-spans["min"]
    if spans["Spanne"].isna().any() or (spans["Spanne"] <= 0).any():
        raise ValueError("Wertspanne fehlt/ist null. Für jede Variable müssen Min und Max vorliegen.")
    out["Wertspanne"] = out["Variable"].map(spans["Spanne"])
    out["Fehler_absolut"] = (out["Antwort"]-out["Sollwert"]).abs()
    out["Fehler_%"] = 100*out["Fehler_absolut"]/out["Wertspanne"]
    out["Exakt"] = out["Gültig"] & np.isclose(out["Antwort"], out["Sollwert"], rtol=0, atol=1e-10)
    out["Innerhalb_1%"] = out["Gültig"] & (out["Fehler_%"] <= TOLERANCE_PERCENT + 1e-10)
    out["Exakt"] = out["Exakt"].astype("boolean")
    out["Innerhalb_1%"] = out["Innerhalb_1%"].astype("boolean")
    out.loc[~out["Gültig"], ["Exakt", "Innerhalb_1%"]] = pd.NA
    # Auffällige Fehler pro Aufgabe und Methode markieren; NIEMALS automatisch entfernen.
    out["IQR_auffällig"] = False
    thresholds = []
    for (task, method), grp in out[out["Gültig"]].groupby(["Aufgabe", "Methode"]):
        q1, q3 = grp["Fehler_absolut"].quantile([.25, .75]); iqr = q3-q1
        upper = q3 + 1.5*iqr
        # Bei IQR=0 ist die Markierung wenig aussagekräftig: nur dokumentieren.
        applicable = len(grp) >= 4 and iqr > 0
        out.loc[grp.index, "IQR_auffällig"] = applicable & (grp["Fehler_absolut"] > upper)
        thresholds.append({"Aufgabe": task, "Methode": method, "n": len(grp), "Q1": q1,
                           "Q3": q3, "IQR": iqr, "Obere_Grenze": upper,
                           "IQR_anwendbar": applicable,
                           "Auffällige": int(out.loc[grp.index, "IQR_auffällig"].sum())})
    return out, pd.DataFrame(thresholds), spans.reset_index()
def rank_table(wide):
    """Gleichstände mit Mittelrängen; Erstplätze bei Gleichstand anteilig vergeben."""
    ranks = wide.rank(axis=1, method="average", ascending=True)
    wins = wide.eq(wide.min(axis=1), axis=0).div(wide.eq(wide.min(axis=1), axis=0).sum(axis=1), axis=0)
    ties = wide.apply(lambda row: row.map(row.value_counts()).gt(1), axis=1)
    rows = []
    for method in METHODS:
        vals = ranks[method]
        rows.append({"Methode": method, "n": len(wide), "Mittlerer_Rang": vals.mean(),
                     "Anteil_Erstplatz_%": wins[method].mean()*100,
                     "Erstplatz_gewichtet": wins[method].sum(),
                     "Rang_1": int((vals == 1).sum()), "Rang_2": int((vals == 2).sum()),
                     "Rang_3": int((vals == 3).sum()), "An Gleichständen beteiligt": int(ties[method].sum())})
    return pd.DataFrame(rows), ranks, wins
def repeated_anova(wide, title):
    """Einfaktorielle RM-ANOVA inkl. Mauchly, GG und HF (explorativ).
    Vollständige Methodentripel; Korrekturen ändern nur die F-Freiheitsgrade.
    Mauchly wird bei singulärer Kontrast-Kovarianz nicht ausgewertet.
    """
    complete = wide.reindex(columns=METHODS).dropna()
    n, k = complete.shape
    result = {"Vergleich": title, "Test": "ANOVA (Messwiederholung)", "n": n,
              "Hinweis": "Explorativ; GG-korrigiertes p bevorzugen. Mauchly bei kleinem n wenig trennscharf."}
    if n < 3:
        result["Hinweis"] = "Zu wenige vollständige Personen"
        return result
    x = complete.to_numpy(dtype=float)
    grand = x.mean()
    ss_method = float(n * np.square(x.mean(axis=0) - grand).sum())
    ss_person = float(k * np.square(x.mean(axis=1) - grand).sum())
    ss_total = float(np.square(x - grand).sum())
    ss_error = max(0., ss_total - ss_method - ss_person)
    df_method, df_error = k - 1, (n - 1) * (k - 1)
    ms_method = ss_method / df_method
    ms_error = ss_error / df_error
    result.update({"SS_Methode": ss_method, "SS_Person": ss_person,
                   "SS_Residual": ss_error, "SS_Gesamt": ss_total,
                   "MS_Methode": ms_method, "MS_Residual": ms_error,
                   "df1": df_method, "df2": df_error})
    if ms_error <= 1e-12:
        result["Hinweis"] = "Residualvarianz null; F-Test nicht bestimmbar"
        return result
    f = ms_method / ms_error
    result.update({"Statistik": float(f), "p": float(stats.f.sf(f, df_method, df_error)),
                   "Eta2_partiell": float(ss_method / (ss_method + ss_error))})
    # Orthogonale Kontraste entfernen die personenspezifische Niveauverschiebung.
    # Helmert-Kontraste, für k=3 äquivalent zu einer beliebigen orthonormalen
    # Basis des Unterraums mit Kontrastsummen null.
    from scipy.linalg import helmert
    contrast = helmert(k, full=False)
    covariance = np.cov(x @ contrast.T, rowvar=False, ddof=1)
    eigenvalues = np.linalg.eigvalsh(np.atleast_2d(covariance))
    d = k - 1
    trace = float(eigenvalues.sum())
    trace_sq = float(np.square(eigenvalues).sum())
    if trace <= 1e-12 or trace_sq <= 1e-24:
        result["Hinweis"] = "Kontrastkovarianz degeneriert; Korrekturen nicht bestimmbar"
        return result
    eps_gg = float(np.clip(trace**2 / (d * trace_sq), 1.0 / d, 1.0))
    df1_gg, df2_gg = df_method * eps_gg, df_error * eps_gg
    result.update({"Epsilon_GG": eps_gg, "df1_GG": df1_gg, "df2_GG": df2_gg,
                   "p_GG": float(stats.f.sf(f, df1_gg, df2_gg))})
    # Huynh-Feldt-Schätzer; auf [1/(k-1), 1] begrenzt.
    denom = d * (n - 1 - d * eps_gg)
    if denom > 0:
        eps_hf = float(np.clip((n * d * eps_gg - 2) / denom, 1.0 / d, 1.0))
        result.update({"Epsilon_HF": eps_hf, "df1_HF": df_method * eps_hf,
                       "df2_HF": df_error * eps_hf,
                       "p_HF": float(stats.f.sf(f, df_method * eps_hf, df_error * eps_hf))})
    # Mauchlys W: geometrisches / arithmetisches Mittel der Eigenwerte
    # (als Potenz d); Bartlett-Chi-Quadrat-Approximation.
    if np.min(eigenvalues) > 1e-12:
        mauchly_w = float(np.prod(eigenvalues) / (trace / d)**d)
        mauchly_w = float(np.clip(mauchly_w, np.finfo(float).tiny, 1.0))
        mauchly_df = d * (d + 1) / 2 - 1
        correction = 1 - (2 * d*d + d + 2) / (6 * d * (n - 1))
        if correction > 0:
            mauchly_chi2 = float(-(n - 1) * correction * np.log(mauchly_w))
            result.update({"Mauchly_W": mauchly_w, "Mauchly_Chi2": mauchly_chi2,
                           "Mauchly_df": mauchly_df,
                           "Mauchly_p": float(stats.chi2.sf(mauchly_chi2, mauchly_df))})
    else:
        result["Hinweis"] += " Mauchly nicht bestimmbar (singuläre Kovarianz)."
    return result
def order_test_columns(frame):
    """Friedman-Entscheidung direkt hinter Signifikanz; ANOVA-Kennzahlen danach."""
    preferred = ['Vergleich', 'Test', 'n', 'Statistik', 'p', 'Kendall_W',
                 'Signifikanz (p < 0.05)', 'p_Holm_4_primäre_Friedman',
                 'Primär_Friedman_signifikant', 'Hinweis']
    return frame.reindex(columns=[c for c in preferred if c in frame.columns] +
                         [c for c in frame.columns if c not in preferred])
def combined_tests(wide, title):
    """Friedman (ggf. Wilcoxon/Holm), direkt darunter Messwiederholungs-ANOVA."""
    frame = pd.DataFrame([*friedman_comparison(wide, title), repeated_anova(wide, title)])
    # ANOVA: primär den GG-korrigierten p-Wert beurteilen, nicht das unkorrigierte p.
    mask = frame['Test'].eq('ANOVA (Messwiederholung)')
    for idx in frame.index[mask]:
        p_gg = frame.at[idx, 'p_GG'] if 'p_GG' in frame.columns else np.nan
        frame.at[idx, 'Signifikanz (p < 0.05)'] = (
            'Signifikant (GG)' if pd.notna(p_gg) and p_gg < .05 else
            'Nicht signifikant (GG)' if pd.notna(p_gg) else
            'Nicht bestimmbar (GG)')
    return order_test_columns(frame)
def friedman_comparison(wide, title):
    """Gepaarter Friedman-Test; paarweise Wilcoxon nur bei p<.05, Holm-Korrektur."""
    if len(wide) < 3:
        return [{"Vergleich": title, "Test": "Friedman", "n": len(wide), "Hinweis": "Zu wenige vollständige Personen"}]
    if (wide.nunique(axis=1) == 1).all():
        return [{"Vergleich": title, "Test": "Friedman", "n": len(wide), "Hinweis": "Alle Methoden identisch"}]
    stat, p = stats.friedmanchisquare(*(wide[m] for m in METHODS))
    result = [{"Vergleich": title, "Test": "Friedman", "n": len(wide), "Statistik": stat, "p": p,
               "Kendall_W": stat/(len(wide)*(len(METHODS)-1)), "Signifikanz (p < 0.05)": "Signifikant" if p < 0.05 else "Nicht signifikant", "Hinweis": ""}]
    if p < .05:
        pairs = [(METHODS[0], METHODS[1]), (METHODS[0], METHODS[2]), (METHODS[1], METHODS[2])]
        tests = []
        for a, b in pairs:
            diff = wide[a]-wide[b]
            if np.allclose(diff, 0):
                tests.append((a, b, 0., 1.))
            else:
                w, wp = stats.wilcoxon(wide[a], wide[b], zero_method="wilcox", alternative="two-sided")
                tests.append((a, b, w, wp))
        # Holm: sortierte p-Werte schrittweise adjustieren, monoton halten.
        ordered = sorted(range(3), key=lambda i: tests[i][3]); adj = [None]*3; previous = 0.
        for k, idx in enumerate(ordered):
            previous = max(previous, min(1., tests[idx][3]*(3-k)))
            adj[idx] = previous
        for (a,b,w,wp), ap in zip(tests, adj):
            result.append({"Vergleich": title, "Test": f"Wilcoxon: {a} / {b}", "n": len(wide),
                           "Statistik": w, "p": wp, "p_Holm": ap, **paired_effects(wide[a], wide[b]), "Hinweis": "Paarweise, Holm innerhalb der Aufgabe"})
    return result

def holm_adjust(pvalues):
    """Holm-adjustierte p-Werte; fehlende Tests bleiben NaN."""
    result = [np.nan] * len(pvalues)
    indices = sorted((i for i, p in enumerate(pvalues) if pd.notna(p)), key=lambda i: pvalues[i])
    previous = 0.0
    for rank, i in enumerate(indices):
        previous = max(previous, min(1.0, float(pvalues[i]) * (len(indices) - rank)))
        result[i] = previous
    return result
def paired_effects(a, b):
    """Differenz A-B; rank-biseriale Korrelation auf nichtnull Differenzen."""
    d = (pd.Series(a).reset_index(drop=True) - pd.Series(b).reset_index(drop=True)).dropna()
    nonzero = d[~np.isclose(d, 0, atol=1e-12)]
    if nonzero.empty:
        rbc = 0.0 if len(d) else np.nan
    else:
        ranks = stats.rankdata(abs(nonzero.to_numpy(float)))
        plus = ranks[nonzero.to_numpy(float) > 0].sum()
        minus = ranks[nonzero.to_numpy(float) < 0].sum()
        rbc = float((plus - minus) / (plus + minus))
    return {'Median_Differenz_A_minus_B': float(d.median()) if len(d) else np.nan,
            'Mittel_Differenz_A_minus_B': float(d.mean()) if len(d) else np.nan,
            'Rangbiserial_r': rbc, 'n_Nichtnull_Differenzen': len(nonzero)}
def task_comparisons(valid):
    """Sekundär: gepaarte Max/Min-Tests und explorative Differenz-der-Differenzen."""
    rows, individual = [], []
    for label, maximum, minimum in [
        ('Unsicherheit', 'Niederschlag_Uncertainty_Maxima', 'Niederschlag_Uncertainty_Minima'),
        ('Mittelwert', 'Luftdruck_Value_Maxima', 'Luftdruck_Value_Minima')]:
        part = valid[valid.Aufgabe.isin([maximum, minimum])]
        wide = part.pivot(index='Teilnehmer', columns=['Aufgabe', 'Methode'], values='Fehler_%')
        changes = pd.DataFrame(index=wide.index)
        for method in METHODS:
            if (maximum, method) not in wide.columns or (minimum, method) not in wide.columns:
                continue
            pair = wide[[(maximum, method), (minimum, method)]].dropna()
            if pair.empty:
                continue
            a, b = pair.iloc[:, 0], pair.iloc[:, 1]
            changes.loc[pair.index, method] = a - b
            if len(pair) >= 2 and not np.allclose(a, b):
                statistic, pvalue = stats.wilcoxon(a, b, zero_method='wilcox', alternative='two-sided')
            elif len(pair) >= 2:
                statistic, pvalue = 0.0, 1.0
            else:
                statistic, pvalue = np.nan, np.nan
            rows.append({'Bereich': label, 'Test': 'Max gegen Min (Wilcoxon)', 'Methode': method,
                         'n': len(pair), 'Statistik': statistic, 'p': pvalue,
                         **paired_effects(a, b)})
        complete = changes.reindex(columns=METHODS).dropna()
        if len(complete) >= 3 and not (complete.nunique(axis=1) == 1).all():
            stat, pval = stats.friedmanchisquare(*(complete[m] for m in METHODS))
            w = stat / (len(complete) * (len(METHODS) - 1))
        elif len(complete) >= 3:
            stat, pval, w = 0.0, 1.0, 0.0
        else:
            stat, pval, w = np.nan, np.nan, np.nan
        rows.append({'Bereich': label, 'Test': 'Methodenvergleich der Max-Min-Differenzen (Friedman)',
                     'Methode': 'alle drei', 'n': len(complete), 'Statistik': stat,
                     'p': pval, 'Kendall_W': w,
                     'Hinweis': 'Explorativer Rangvergleich der Differenzen; kein vollständiger faktorieller Interaktionstest'})
        for participant, vals in changes.iterrows():
            for method in METHODS:
                if pd.notna(vals.get(method, np.nan)):
                    individual.append({'Bereich': label, 'Teilnehmer': participant, 'Methode': method,
                                       'Fehler_Max_minus_Min_Prozentpunkte': vals[method]})
    tests = pd.DataFrame(rows)
    if not tests.empty:
        tests['p_Holm_6_Wilcoxon'] = np.nan
        mask = tests.Test.eq('Max gegen Min (Wilcoxon)')
        tests.loc[mask, 'p_Holm_6_Wilcoxon'] = holm_adjust(tests.loc[mask, 'p'].tolist())
        tests['p_Holm_2_Friedman'] = np.nan
        mask = tests.Test.str.startswith('Methodenvergleich')
        tests.loc[mask, 'p_Holm_2_Friedman'] = holm_adjust(tests.loc[mask, 'p'].tolist())
        wilcoxon = tests.Test.eq('Max gegen Min (Wilcoxon)')
        def richtung(median):
            if pd.isna(median):
                return 'Nicht bestimmbar'
            if median > 0:
                return 'Maximum ungenauer'
            if median < 0:
                return 'Minimum ungenauer'
            return 'Kein Unterschied im Median'
        tests.loc[wilcoxon, 'Richtung_der_Differenz'] = (
            tests.loc[wilcoxon, 'Median_Differenz_A_minus_B'].map(richtung)
        )
        tests.loc[wilcoxon, 'Statistisches_Ergebnis'] = tests.loc[wilcoxon, 'p_Holm_6_Wilcoxon'].map(
            lambda p: 'Nicht berechenbar' if pd.isna(p) else ('Signifikant' if p < 0.05 else 'Nicht signifikant')
        )
    return tests, pd.DataFrame(individual)
def grouped_plot(table, value, ylabel, folder, filename, scientific=False):
    """Vier Aufgabentypen, jeweils drei Methoden nebeneinander; Legende fest rechts."""
    fig, ax = plt.subplots(figsize=(10, 5.5)); shown = [t for t in TASKS if t in table.Aufgabe.values]
    x = np.arange(len(shown)); width = .25
    for i, method in enumerate(METHODS):
        vals = [table.loc[(table.Aufgabe == t)&(table.Methode == method), value].iloc[0]
                if ((table.Aufgabe == t)&(table.Methode == method)).any() else np.nan for t in shown]
        ax.bar(x+(i-1)*width, vals, width, color=COLORS[method], label=METHOD_LABELS[method])
    ax.set_xticks(x, [TASK_LABELS[t] for t in shown], rotation=15, ha="right")
    ax.set_ylabel(ylabel); ax.set_ylim(bottom=0)
    if scientific:
        # Wissenschaftliche Notation als ×10ⁿ statt der Schreibweise 1e-n.
        ax.ticklabel_format(axis="y", style="sci", scilimits=(0, 0), useMathText=True)
    ax.legend(loc="center left", bbox_to_anchor=(1.02,.5), borderaxespad=0)
    ax.spines[["top","right"]].set_visible(False)
    save_plot(fig, folder, filename, legend=True)
def no_flag_summary(noflag, original_people):
    """Explorative Sensitivität: gleiche ursprüngliche Personengruppe, markierte Werte entfernt.
    Personen mit danach unvollständigen Aufgaben werden nicht mit wechselnden Aufgabenmitteln vermischt.
    """
    w=noflag[noflag.Teilnehmer.isin(original_people)].pivot(index=["Teilnehmer","Aufgabe"],columns="Methode",values="Fehler_%").reindex(columns=METHODS).dropna()
    n=w.reset_index().groupby("Teilnehmer")["Aufgabe"].nunique() if len(w) else pd.Series(dtype=int)
    ids=n[n==len(TASKS)].index
    w=w.loc[w.index.get_level_values("Teilnehmer").isin(ids)]
    w=w.groupby(level="Teilnehmer").mean() if len(w) else pd.DataFrame(columns=METHODS)
    return pd.DataFrame([{"Methode":m,"n_Personen_ohne_markierte":len(w),
                          "Mittelwert_ohne_markierte":w[m].mean() if len(w) else np.nan} for m in METHODS])
# Lesbare Skalierung: Niederschlag in 10^-5 mm/Tag, Luftdruck unverändert in Pa.
SCALE = {'Niederschlag_Uncertainty': (100000., '10⁻⁵ mm/Tag'),
         'Luftdruck_Value': (1., 'Pa')}
def _excel_text(value):
    """Konvertiert versehentliche literale Zeilenumbrüche in echte Umbrüche."""
    if isinstance(value, str):
        return value.replace('\\n', '\n').replace('/n', '\n')
    return value
def _estimated_lines(value, width):
    """Schätzt die nötige Zeilenhöhe einschließlich automatischer Umbrüche."""
    text = str(_excel_text(value) if value is not None else '')
    usable = max(8, width - 3)
    return sum(max(1, (len(part) + usable - 1) // usable) for part in text.split('\n'))
def write_section(writer, sheetname, sections, note):
    """Vollbreite Abschnittsbalken, lesbare Spalten und variable Zeilenhöhen."""
    sheet = writer.book.create_sheet(sheetname)
    sheet.sheet_view.showGridLines = False
    ncols = max([2] + [len(frame.columns) for _, frame, _ in sections if frame is not None and len(frame.columns)])
    # Spaltenbreiten anhand von Überschriften und tatsächlichen Zellinhalten bestimmen.
    widths = [18.0] * ncols
    for _, frame, _ in sections:
        if frame is None or not len(frame.columns):
            continue
        for j, column in enumerate(frame.columns):
            samples = [str(column)] + [str(v) for v in frame.iloc[:200, j] if pd.notna(v)]
            longest = max((max(map(len, str(_excel_text(v)).split('\n'))) for v in samples), default=16)
            widths[j] = max(widths[j], min(65.0, longest + 3.0))
    # J bleibt in den vier Aufgabenblättern kompakt; lange Hinweise umbrechen.
    if sheetname in {TASK_LABELS[t].replace(': ', '_')[:31] for t in TASKS} and ncols >= 10:
        widths[9] = 22.0
    if sheetname == '00_Lesehilfe':
        widths[0] = max(widths[0], 31)
        widths[1] = 92
    for j, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(j)].width = width
    # Merged-Zellen erhalten über die gesamte tatsächlich verwendete Tabellenbreite Füllung.
    def banner(row, message, color, font_color='FFFFFF', bold=True):
        for j in range(1, ncols + 1):
            sheet.cell(row, j).fill = PatternFill('solid', fgColor=color)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell = sheet.cell(row, 1, _excel_text(message))
        cell.font = Font(bold=bold, color=font_color, size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        available = sum(widths) * .85
        sheet.row_dimensions[row].height = max(29, min(150, 17 * _estimated_lines(message, available) + 12))
    banner(1, note, '23415C')
    row = 3
    for title, frame, description in sections:
        banner(row, title, '3676B6')
        row += 1
        if description:
            banner(row, description, 'EAF0F6', font_color='20354A', bold=False)
            row += 1
        if frame is not None and len(frame.columns):
            for j, name in enumerate(frame.columns, 1):
                cell = sheet.cell(row, j, _excel_text(str(name)))
                cell.font = Font(bold=True, color='20354A')
                cell.fill = PatternFill('solid', fgColor='D7E5F1')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.row_dimensions[row].height = max(29, min(95, 16 * max(_estimated_lines(name, widths[j]) for j, name in enumerate(frame.columns)) + 8))
            row += 1
            for values in frame.itertuples(index=False, name=None):
                lines = 1
                for j, value in enumerate(values, 1):
                    if pd.isna(value): value = None
                    if isinstance(value, np.generic): value = value.item()
                    value = _excel_text(value)
                    cell = sheet.cell(row, j, value)
                    cell.alignment = Alignment(horizontal='center' if isinstance(value, (int, float, np.integer, np.floating)) and not isinstance(value, (bool, np.bool_)) else 'left', vertical='top', wrap_text=True)
                    if isinstance(value, (float, np.floating)):
                        col_name = str(frame.columns[j-1])
                        # Einzelne Likert-Antworten sind ganze Skalenwerte; Mittelwerte nicht.
                        if col_name == 'Likert_Response' or col_name.startswith('Likert_') and not col_name.startswith('Likert_Mittel_'):
                            cell.number_format = '0'
                        else:
                            cell.number_format = '#,##0.0000;[Red]-#,##0.0000'
                    if value is not None:
                        lines = max(lines, _estimated_lines(value, widths[j-1]))
                sheet.row_dimensions[row].height = max(19, min(420, 15 * lines + 5))
                row += 1
        row += 2
    sheet.freeze_panes = None
def boxplot_groups(data, group_tasks, output, name, ylabel, metric='Fehler_%', title=None):
    """Standard-Boxplots mit Whiskern 1,5 × IQR, nur Ausreißerpunkte und n.
    n zählt Personen mit einem gültigen Wert pro Methode. Bei unterschiedlichen
    Stichprobengrößen werden die methodenspezifischen n explizit ausgewiesen.
    """
    fig, ax = plt.subplots(figsize=(7.5, 5.2))
    selected = data[data.Aufgabe.isin(group_tasks)].copy()
    selected = selected.dropna(subset=[metric])
    counts = {}
    for i, method in enumerate(METHODS, 1):
        rows = selected[selected.Methode.eq(method)].sort_values('Teilnehmer')
        if rows.Teilnehmer.duplicated().any():
            raise ValueError(f'Mehrere Werte pro Person/Methode im Boxplot {name}: {method}')
        values = rows[metric].to_numpy(float)
        counts[method] = len(values)
        if not len(values):
            continue
        bp = ax.boxplot([values], positions=[i], widths=.52, patch_artist=True,
                        whis=1.5, showmeans=False, showfliers=True,
                        flierprops={'marker': 'o', 'markerfacecolor': COLORS[method],
                                    'markeredgecolor': 'white', 'markersize': 5, 'alpha': .85})
        bp['boxes'][0].set(facecolor=COLORS[method], alpha=.3, edgecolor=COLORS[method])
        for item in bp['medians']:
            item.set(color=COLORS[method], linewidth=2)

    # Die Annotation zählt die gültigen Werte in der Boxplot-Berechnung und nicht
    # nur die als Ausreißer sichtbaren Punkte.
    if len(set(counts.values())) == 1:
        n_label = f'n = {next(iter(counts.values()))}'
    else:
        n_label = 'n: ' + ', '.join(f'{METHOD_LABELS[m]} {counts[m]}' for m in METHODS)
    # Kleine n-Angabe unterhalb der Achse, außerhalb der Boxplots/Punkte.
    ax.text(0.0, -0.17, n_label, transform=ax.transAxes, ha='left', va='top',
            fontsize=7.5, color='#555555', clip_on=False)
    ax.set_xticks(range(1, 4), [METHOD_LABELS[m] for m in METHODS])
    ax.set_ylabel(ylabel)
    if title:
        ax.set_title(title)
    ax.spines[['top', 'right']].set_visible(False)
    save_plot(fig, output, name)
def paired_group_means(valid, tasks, metric, label):
    """Pro Person und Methode mitteln; nur vollständige Tripel aller Aufgaben."""
    subset = valid[valid.Aufgabe.isin(tasks)].dropna(subset=[metric])
    wide = subset.pivot(index=['Teilnehmer', 'Aufgabe'], columns='Methode',
                        values=metric).reindex(columns=METHODS).dropna()
    if wide.empty:
        return pd.DataFrame(columns=['Teilnehmer', 'Methode', metric, 'Aufgabe'])
    counts = wide.reset_index().groupby('Teilnehmer').Aufgabe.nunique()
    people = counts[counts.eq(len(tasks))].index
    wide = wide[wide.index.get_level_values('Teilnehmer').isin(people)]
    mean = wide.groupby(level='Teilnehmer').mean()
    out = mean.reset_index().melt(id_vars='Teilnehmer', var_name='Methode', value_name=metric)
    out['Aufgabe'] = label
    return out
# Verständliche Kurzbeschreibungen in den jeweiligen Statistikabschnitten.
TEST_EXPLANATION_PRIMARY = (
    'Fragestellung: Unterscheiden sich VSUP, ScaledGlyph und IsoGlyph beim normalisierten Fehler dieser Aufgabe? '
    'Friedman vergleicht die drei Methoden anhand der Ränge innerhalb derselben Personen (nur vollständige Methodentripel). '
    'Ein kleiner p-Wert spricht gegen gleiche Rangverteilungen; Kendall W (0 bis 1) beschreibt die Stärke der Rangübereinstimmung bzw. Methodendifferenz. '
    'Primäre Entscheidung: p_Holm_4_primäre_Friedman < 0,05 (Korrektur über vier Aufgaben). '
    'Nur danach werden paarweise Wilcoxon-Vergleiche als bestätigende Folgetests ausgewiesen; deren p_Holm korrigiert drei Methodenpaare innerhalb der Aufgabe. '
    'Die ANOVA mit Messwiederholung prüft ergänzend Mittelwertunterschiede; ihre Signifikanzspalte bewertet p_GG (Greenhouse-Geisser-korrigiert), nicht das unbereinigte p. Mit Mauchly-Test und Huynh-Feldt-Korrektur; explorativ.'
)
TEST_EXPLANATION_OVERALL = (
    'Ergänzende Gesamtanalyse: Pro Person und Methode wird zunächst der Mittelwert der vier normalisierten Aufgabenfehler gebildet '
    '(nur Personen mit allen zwölf gültigen Antworten). Friedman prüft Unterschiede zwischen den drei Methoden anhand persönlicher Ränge; '
    'Kendall W beschreibt die Effektstärke. Bei unbereinigtem Friedman-p < 0,05 folgen drei paarweise Wilcoxon-Tests mit Holm-Korrektur. '
    'Die ANOVA mit Messwiederholung vergleicht ergänzend die Methodenmittelwerte; die Signifikanzspalte nutzt p_GG (explorativ). '
    'Der Gesamtvergleich ersetzt nicht die vier primären Aufgabenanalysen.'
)
TEST_EXPLANATION_TIME = (
    'Sekundäre Analyse der Bearbeitungszeit: Pro Person und Methode wird die mittlere Zeit über vier Aufgaben gebildet '
    '(nur vollständige Zeitangaben). Friedman prüft Rangunterschiede zwischen den Methoden; Kendall W beschreibt deren Stärke. '
    'Bei p < 0,05 folgen paarweise Wilcoxon-Tests mit Holm-Korrektur. Die ANOVA vergleicht ergänzend die mittleren Zeiten '
    'mit Greenhouse-Geisser- und Huynh-Feldt-Korrektur; die Signifikanzspalte nutzt p_GG (explorativ).'
)
def main():
    base = Path(__file__).resolve().parent
    source = base / CSV_FILENAME
    if not source.exists(): raise FileNotFoundError(f'CSV nicht gefunden: {source}')
    output = base / OUTPUT_FOLDER
    output.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(source, low_memory=False, dtype={'participantId': str})
    required = {'participantId', 'status', 'trialId', 'answer', 'correctAnswer', 'duration'}
    if required - set(df.columns): raise ValueError(f'Fehlende Spalten: {required-set(df.columns)}')
    info = df.groupby('participantId').status.agg(
        completed=lambda s: s.eq('completed').any(),
        rejected=lambda s: s.eq('rejected').any())
    keep = info.index[info.completed & ~info.rejected]
    df = df[df.participantId.isin(keep)].copy()
    trials, thresholds, spans = parse_trials(df)
    trials['Skalierungsfaktor'] = trials.Variable.map(lambda v: SCALE[v][0])
    trials['Einheit_skaliert'] = trials.Variable.map(lambda v: SCALE[v][1])
    for source_col, target_col in [('Antwort', 'Antwort_skaliert'),
                                   ('Sollwert', 'Sollwert_skaliert'),
                                   ('Wertspanne', 'Wertspanne_skaliert'),
                                   ('Fehler_absolut', 'Fehler_absolut_skaliert')]:
        trials[target_col] = trials[source_col] * trials.Skalierungsfaktor
    valid = trials[trials.Gültig].copy()
    triples = valid.pivot(index=['Teilnehmer','Aufgabe'], columns='Methode', values='Fehler_%').reindex(columns=METHODS).dropna()
    counts = triples.reset_index().groupby('Teilnehmer').Aufgabe.nunique()
    complete_ids = counts[counts.eq(len(TASKS))].index
    overall_data = valid[valid.Teilnehmer.isin(complete_ids)].copy()
    overall = overall_data.groupby(['Teilnehmer', 'Methode'])['Fehler_%'].mean().unstack().reindex(columns=METHODS).dropna()
    overall_ranks, personal_ranks, _ = rank_table(overall) if len(overall) else (pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
    if len(personal_ranks):
        likert_means = valid[valid.Teilnehmer.isin(overall.index)].groupby(['Teilnehmer', 'Methode'])['Likert_Response'].mean().unstack().reindex(columns=METHODS)
        enriched = pd.DataFrame(index=personal_ranks.index)
        for method in METHODS:
            enriched[method] = personal_ranks[method]
            enriched[f'Likert_Mittel_{METHOD_LABELS[method]}'] = likert_means.reindex(personal_ranks.index)[method]
        personal_ranks = enriched
    overall_desc = pd.DataFrame([{'Methode': m, **describe(overall[m])} for m in METHODS])
    tests_total = combined_tests(overall, 'Gesamt: normalisierte Abweichung')
    task_tests, task_differences = task_comparisons(valid)
    primary_p = {}
    for task in TASKS:
        w_task = triples.xs(task, level='Aufgabe') if task in triples.index.get_level_values('Aufgabe') else pd.DataFrame(columns=METHODS)
        result = friedman_comparison(w_task, task)
        primary_p[task] = result[0].get('p', np.nan)
    primary_adjusted = dict(zip(TASKS, holm_adjust([primary_p[t] for t in TASKS])))
    # Aufgabenblätter: zuerst alle Antworten, anschließend Deskription, Ränge und Tests.
    sheets = []
    for task in TASKS:
        part = trials[trials.Aufgabe.eq(task)].copy()
        good = part[part.Gültig]
        rows = []
        for method in METHODS:
            g = good[good.Methode.eq(method)]
            d1 = describe(g.Fehler_absolut_skaliert)
            d2 = describe(g['Fehler_%'])
            dt = describe(g.Zeit_s)
            rows.append({'Methode': method, 'n gültig': len(g), 'n ungültig': len(part[part.Methode.eq(method)])-len(g),
                         'Exakt': int(g.Exakt.sum()), 'Innerhalb 1 %': int(g['Innerhalb_1%'].sum()),
                         'Fehler skaliert Mittel': d1['Mittelwert'], 'Fehler skaliert Median': d1['Median'],
                         'Fehler skaliert SD': d1['Standardabweichung'], 'Fehler % Mittel': d2['Mittelwert'],
                         'Fehler % Median': d2['Median'], 'Fehler % SD': d2['Standardabweichung'],
                         'Fehler % Q1': d2['Q1'], 'Fehler % Q3': d2['Q3'], 'Fehler % IQR': d2['IQR'],
                         'Zeit s Mittel': dt['Mittelwert'], 'Zeit s Median': dt['Median']})
        desc = pd.DataFrame(rows)
        w = triples.xs(task, level='Aufgabe') if task in triples.index.get_level_values('Aufgabe') else pd.DataFrame(columns=METHODS)
        ranks, person_ranks, _ = rank_table(w) if len(w) else (pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
        if len(person_ranks):
            task_likert = part.pivot(index='Teilnehmer', columns='Methode', values='Likert_Response')
            enriched = pd.DataFrame(index=person_ranks.index)
            for method in METHODS:
                enriched[method] = person_ranks[method]
                enriched[f'Likert_{METHOD_LABELS[method]}'] = task_likert.reindex(person_ranks.index)[method]
            person_ranks = enriched
        stat = combined_tests(w, task)
        stat['p_Holm_4_primäre_Friedman'] = np.nan
        stat.loc[stat.Test.eq('Friedman'), 'p_Holm_4_primäre_Friedman'] = primary_adjusted[task]
        stat['Primär_Friedman_signifikant'] = ''
        if pd.notna(primary_adjusted[task]):
            stat.loc[stat.Test.eq('Friedman'), 'Primär_Friedman_signifikant'] = ('Ja' if primary_adjusted[task] < .05 else 'Nein')
        stat = order_test_columns(stat)
        # Paarvergleiche nur dann als konfirmatorisch interpretieren, wenn
        # der über vier Aufgaben korrigierte primäre Friedman-Test signifikant ist.
        if not (pd.notna(primary_adjusted[task]) and primary_adjusted[task] < .05):
            stat = stat.loc[~stat.Test.str.startswith('Wilcoxon')].copy()
        scale = SCALE[part.Variable.iloc[0]]
        note = f'{TASK_LABELS[task]} | skalierte absolute Fehler in {scale[1]}; normalisierte Fehler in % der vollständigen Variablenspanne.'
        answer_columns = ['Teilnehmer','Methode','Antwort_skaliert','Sollwert_skaliert','Wertspanne_skaliert',
                          'Skalierungsfaktor','Einheit_skaliert','Fehler_absolut_skaliert','Fehler_%','Likert_Response','Gültig','Zeit_s','IQR_auffällig']
        sheets.append((TASK_LABELS[task].replace(': ', '_')[:31], [
            ('A. Antworten und Einzelberechnungen', part[answer_columns].sort_values(['Teilnehmer','Methode']),
             f'Absoluter Fehler (skaliert) = |Antwort − Sollwert| × {scale[0]:g} ({scale[1]}); Fehler (%) = 100 × |Antwort − Sollwert| / Wertspanne. Faktor ist in der Tabelle ausgewiesen.'),
            ('B. Deskriptive Kennzahlen', desc, 'Mittel = Summe / n; Median = mittlerer sortierter Wert; SD = Stichprobenstandardabweichung (ddof=1); Q1/Q3 = 25./75. Perzentil; IQR = Q3 − Q1.'),
            ('C. Persönliche Ränge', person_ranks.reset_index() if len(person_ranks) else pd.DataFrame(), 'Nur vollständige Methodentripel; kleinster Fehler erhält Rang 1; Gleichstände erhalten Mittelränge. Neben dem Rang steht die Likert-Sicherheit (1–7) je Methode.'),
            ('D. Zusammenfassung der Ränge', ranks, 'Mittlerer Rang = Summe der persönlichen Mittelränge / n; Erstplatz_gewichtet = Summe(1 / Anzahl geteilter Erstplätze), sofern Methode erstplatziert, sonst 0; Anteil_Erstplatz_% = 100 × Erstplatz_gewichtet / n. Rang_1/2/3 zählen exakte Ränge; Gleichstandsbeteiligung zählt auch geteilte andere Ränge.'),
            ('E. Statistische Tests', stat, TEST_EXPLANATION_PRIMARY)], note))
        # 01: Je Teilaufgabe normalisierte Fehler und Zeiten.
        for suffix, metric, ylabel in [
            ('Prozent', 'Fehler_%', 'Normalisierte Abweichung (%)'),
            ('Zeit', 'Zeit_s', 'Bearbeitungszeit (s)'),
        ]:
            boxplot_groups(good, [task], output,
                           f'01_{TASK_LABELS[task].replace(": ", "_")}_{suffix}',
                           ylabel, metric=metric, title=TASK_LABELS[task])
    # Einzelantworten: getrennte Datei, originale + skalierte Werte.
    with pd.ExcelWriter(output / '00_Identifikation_Einzelantworten.xlsx', engine='openpyxl') as writer:
        trials.to_excel(writer, sheet_name='Einzelantworten', index=False)
        sh = writer.sheets['Einzelantworten']
        sh.freeze_panes = None
        sh.auto_filter.ref = sh.dimensions
        sh.sheet_view.showGridLines = False
        for cell in sh[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='23415C')
        for j, name in enumerate(trials.columns, 1):
            col = sh.column_dimensions[get_column_letter(j)]
            samples = [str(name)] + [str(v) for v in trials.iloc[:200, j - 1] if pd.notna(v)]
            col.width = min(72, max(22, max(len(v) for v in samples) + 3))
        for cell in sh[1]: cell.alignment = Alignment(wrap_text=True, vertical='center')
        if 'Likert_Response' in trials.columns:
            j_likert = trials.columns.get_loc('Likert_Response') + 1
            for row in sh.iter_rows(min_row=2, min_col=j_likert, max_col=j_likert):
                row[0].number_format = '0'
        sh.row_dimensions[1].height = 34
        # Zahlen in der separaten Einzelantworten-Datei ebenfalls zentrieren.
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.alignment = Alignment(horizontal='center', vertical='center')
    with pd.ExcelWriter(output / '01_Identifikation_Auswertung.xlsx', engine='openpyxl') as writer:
        intro = pd.DataFrame([
            ['Einschluss', 'Nur completed und nicht rejected; in progress immer ausgeschlossen.'],
            ['Absoluter Fehler', '|Antwort − Sollwert|, Originaleinheit.'],
            ['Skalierter Fehler', 'Absoluter Fehler × Faktor; Niederschlag: Faktor 100000, Luftdruck: Faktor 1.'],
            ['Normalisierter Fehler (%)', '100 × absoluter Fehler / (Maximum − Minimum der jeweiligen Variablenwerte).'],
            ['Mittelwert', 'Arithmetisches Mittel aller gültigen Werte in der jeweiligen Gruppe.'],
            ['Median', '50. Perzentil der gültigen Werte; entfernt keine Ausreißer.'],
            ['SD', 'Stichprobenstandardabweichung, ddof=1.'],
            ['Q1, Q3, IQR', '25. und 75. Perzentil; IQR = Q3 − Q1.'],
            ['Primäre Tests', 'Vier Friedman-Tests der Einzelaufgaben; Holm-Korrektur über die vier p-Werte. Paarweise Wilcoxon/Holm nur nach korrigiert signifikantem Friedman konfirmatorisch.'],
            ['Aufgabenvergleiche', 'Sekundär: sechs gepaarte Max-Min-Wilcoxon-Tests (Holm über sechs) und zwei explorative Friedman-Tests auf Max-Min-Differenzen (Holm über zwei).'],
            ['Gesamtvergleich', 'Nur Personen mit 12 gültigen Antworten: pro Person/Methode Mittel der 4 normalisierten Aufgabenfehler.'],
            ['Einzelaufgaben', 'Deskription: alle gültigen Antworten je Methode; Tests: nur vollständige Methodentripel derselben Aufgabe.'],
            ['Friedman: Fragestellung', 'Für jede Person werden die drei Methoden nach Fehlergröße geordnet. Der Friedman-Test prüft, ob sich ihre Ränge systematisch unterscheiden.'],
            ['Friedman: Datenbasis', 'Nur vollständige Methodentripel je Aufgabe. n = Anzahl Personen, k = 3 Methoden.'],
            ['Friedman: p-Wert', 'Der p-Wert prüft die Nullhypothese ohne systematischen Unterschied zwischen den Methoden. Für die vier primären Aufgaben ist der über vier Tests Holm-korrigierte p-Wert maßgeblich.'],
            ['Friedman: Kendall W', 'Kendall W = χ² / [n × (k − 1)]. Rangbasierte Effektstärke zwischen 0 und 1; größere Werte bedeuten ausgeprägtere Übereinstimmung der persönlichen Methodenrangfolgen.'],
            ['Wilcoxon: Folgetests', 'Nach einem über vier Aufgaben Holm-korrigiert signifikanten Friedman-Test (p < 0,05) werden drei gepaarte Methodenvergleiche mit Wilcoxon durchgeführt. Deren p-Werte werden innerhalb der Aufgabe ebenfalls nach Holm korrigiert.'],
            ['Wilcoxon: Interpretation', 'Median_Differenz_A_minus_B beschreibt die Richtung der Fehlerdifferenz; Rangbiserial_r beschreibt deren rangbasierte Effektstärke. Ein signifikanter Friedman-Test sagt allein noch nicht, welches Methodenpaar sich unterscheidet.'],
            ['ANOVA: Fragestellung', 'Ergänzend wird geprüft, ob sich die Mittelwerte der drei Methoden unterscheiden. Die Messwiederholung berücksichtigt, dass dieselben Personen alle drei Methoden verwendet haben.'],
            ['ANOVA: F und Freiheitsgrade', 'F = MS_Methode / MS_Residual; df1 = k − 1; df2 = (n − 1) × (k − 1). MS bedeutet mittlere Quadratsumme.'],
            ['ANOVA: partielles η²', 'SS_Methode / (SS_Methode + SS_Residual). Beschreibt den Anteil der für den Methodenvergleich relevanten Varianz, der auf die Methode entfällt.'],
            ['ANOVA: Sphärizität', 'Mauchly_W, Mauchly_Chi2, Mauchly_df und Mauchly_p testen die Gleichheit der Varianzen paarweiser Methodendifferenzen. Ein nicht signifikanter Test beweist Sphärizität nicht; bei kleinem n geringe Teststärke.'],
            ['ANOVA: Greenhouse-Geisser', 'Epsilon_GG liegt bei drei Methoden zwischen 0,5 und 1. df1_GG und df2_GG sind die reduzierten Freiheitsgrade; p_GG ist der dazugehörige korrigierte p-Wert. Die F-Prüfgröße bleibt unverändert. GG wird unabhängig vom Mauchly-Ergebnis berichtet.'],
            ['ANOVA: Huynh-Feldt', 'Epsilon_HF, df1_HF, df2_HF und p_HF bilden die alternative Huynh-Feldt-Korrektur ab.'],
            ['ANOVA: Quadratsummen', 'SS_Gesamt = SS_Methode + SS_Person + SS_Residual; MS_Methode = SS_Methode / df1; MS_Residual = SS_Residual / df2.'],
            ['ANOVA: Einordnung', 'Die Spalte Signifikanz (p < 0.05) nutzt bei ANOVA den Greenhouse-Geisser-korrigierten p_GG (Kennzeichnung GG). Bei Friedman basiert sie auf dem unbereinigten p; für die vier primären Aufgaben gilt ausschließlich Primär_Friedman_signifikant anhand p_Holm_4_primäre_Friedman. ANOVA bleibt explorativ.'] ,
            ['Likert_Response', 'Sicherheitsbewertung (1–7) aus der separaten Likert_Response-Zeile derselben Person und Aufgabe/Methode. Fehlende Werte bleiben leer; die statistische Fehleranalyse wird dadurch nicht gefiltert.'],
            ['Likert in persönlichen Rängen', 'Neben jedem Methodenrang steht die Sicherheitsbewertung der Einzelaufgabe; im Gesamtvergleich der Mittelwert der bis zu vier Likert-Antworten je Person und Methode.'],
            ['IQR auffällig', 'Ein Einzelwert ist auffällig, wenn sein absoluter Fehler > Q3 + 1,5 × (Q3 − Q1) innerhalb derselben Aufgabe/Methode liegt. Nur für n ≥ 4 und IQR > 0; kein Ausschluss und kein Beweis für einen Messfehler.'],
            ['Erstplatz_gewichtet', 'Summe anteiliger Erstplätze: alleiniger erster Platz = 1; zwei gleich gute erste Plätze = je 0,5; drei = je 1/3. Anteil_Erstplatz_% = 100 × Summe / n.'],
            ['Boxplots', 'Standard-Boxplot: Box Q1 bis Q3 (50 %), Medianlinie, Whisker bis zum äußersten Wert innerhalb von Q1−1,5×IQR bzw. Q3+1,5×IQR; nur Werte außerhalb als Punkte. Kein Mittelwert. Alle gültigen Werte fließen in die Berechnung ein.'],
        ], columns=['Begriff', 'Berechnung / Bedeutung'])
        write_section(writer, '00_Lesehilfe', [('Berechnungsgrundlagen', intro, '')], 'Hinweise zu Auswertung, Datenbasis und Formeln')
        for name, sections, note in sheets: write_section(writer, name, sections, note)
        person_means = overall.reset_index() if len(overall) else pd.DataFrame()
        write_section(writer, 'Gesamtvergleich', [
            ('A. Persönliche Durchschnittsfehler (%)', person_means, 'Je Person und Methode: Mittel der vier normalisierten Fehler; nur vollständige 12 Antworten.'),
            ('B. Deskriptive Kennzahlen', overall_desc, 'Kennzahlen über persönliche Durchschnittswerte, nicht über ungepaarte Einzelantworten.'),
            ('C. Persönliche Ränge', personal_ranks.reset_index() if len(personal_ranks) else pd.DataFrame(), 'Rang 1 = kleinster persönlicher Durchschnittsfehler; Likert_Mittel = mittlere vorhandene Sicherheitsbewertung über vier Aufgaben je Person/Methode.'),
            ('D. Rangübersicht', overall_ranks, 'Mittlerer Rang = Mittel der individuellen Ränge; Erstplatz_gewichtet = Summe(1/k) bei k gleich guten Erstplatzierten; Erstplatzanteil = 100 × Erstplatz_gewichtet / n.'),
            ('E. Statistische Tests', tests_total, TEST_EXPLANATION_OVERALL),
        ], 'Übergreifender Vergleich: persönliche Mittel der vier normalisierten Teilaufgaben')
        wilcoxon_rows = task_tests[task_tests['Test'].eq('Max gegen Min (Wilcoxon)')].copy()
        friedman_rows = task_tests[task_tests['Test'].str.startswith('Methodenvergleich')].copy()
        wilcoxon_cols = ['Bereich', 'Methode', 'n', 'Statistik', 'p', 'p_Holm_6_Wilcoxon',
                         'Median_Differenz_A_minus_B', 'Mittel_Differenz_A_minus_B',
                         'Rangbiserial_r', 'n_Nichtnull_Differenzen',
                         'Richtung_der_Differenz', 'Statistisches_Ergebnis']
        friedman_cols = ['Bereich', 'Methode', 'n', 'Statistik', 'p', 'p_Holm_2_Friedman', 'Kendall_W', 'Hinweis']
        write_section(writer, 'Aufgabenvergleiche', [
            ('A. Maximum gegen Minimum (Wilcoxon)', wilcoxon_rows.reindex(columns=wilcoxon_cols),
             'Differenz = Fehler Maximum minus Fehler Minimum (Prozentpunkte). Positive Werte: Maximum ungenauer; negative: Minimum ungenauer. Richtung nach Median; Signifikanz nach Holm-Korrektur über sechs Tests (p < 0,05). Die Richtung allein belegt keine Signifikanz.'),
            ('B. Methodenvergleich der Max-Min-Differenzen (Friedman)', friedman_rows.reindex(columns=friedman_cols),
             'Vergleicht die persönlichen Max-Min-Differenzen zwischen den drei Methoden; zwei Friedman-Tests gemeinsam Holm-korrigiert. Explorativ, kein vollständiger faktorieller Interaktionstest.'),
            ('C. Individuelle Max-Min-Differenzen', task_differences,
             'Nur gültige gepaarte Max-Min-Antworten je Person und Methode; keine IQR-Ausschlüsse.')
        ], 'Sekundäre und explorative Vergleiche zwischen Maximum und Minimum')
        write_section(writer, 'IQR_Diagnostik', [('Markierungsgrenzen', thresholds, 'Je Aufgabe und Methode; auffällige Werte bleiben in allen Hauptauswertungen.')], 'Explorative Ausreißerdiagnostik')
        noflag = valid[~valid.IQR_auffällig]
        write_section(writer, 'Sensitivitaet', [('Vergleich ohne IQR-Markierungen', no_flag_summary(noflag, complete_ids), 'Nur zusätzliche Sensitivitätsanalyse; die Hauptanalyse enthält alle gültigen Antworten.')], 'Explorative Sensitivitätsprüfung')
        times = valid.pivot(index=['Teilnehmer','Aufgabe'], columns='Methode', values='Zeit_s').reindex(columns=METHODS).dropna()
        tcounts = times.reset_index().groupby('Teilnehmer').Aufgabe.nunique()
        tpeople = tcounts[tcounts.eq(len(TASKS))].index
        tdata = times.loc[times.index.get_level_values('Teilnehmer').isin(tpeople)].groupby(level='Teilnehmer').mean().reindex(columns=METHODS)
        write_section(writer, 'Bearbeitungszeit', [
            ('A. Persönliche mittlere Zeiten (s)', tdata.reset_index(), 'Nur Personen mit gültigen Zeitwerten in allen vier Methodentripeln.'),
            ('B. Deskriptive Kennzahlen', pd.DataFrame([{'Methode':m, **describe(tdata[m])} for m in METHODS]), 'Kennzahlen über persönliche Mittelzeiten.'),
            ('C. Statistische Tests', combined_tests(tdata, 'Zeit (s)'), TEST_EXPLANATION_TIME)], 'Sekundärer Zeitvergleich')
    # 02: Mittelwert (Min/Max) und Unsicherheit (Min/Max), je Prozent und Zeit.
    # 03: Alle vier Aufgaben, ebenfalls je Prozent und Zeit.
    # Die Zusammenfassungen nutzen vollständige gepaarte Daten pro Metrik.
    for prefix, tasks in [('Mittelwert', TASKS[2:]), ('Unsicherheit', TASKS[:2])]:
        for suffix, metric, ylabel in [
            ('Prozent', 'Fehler_%', 'Mittlere normalisierte Abweichung (%)'),
            ('Zeit', 'Zeit_s', 'Mittlere Bearbeitungszeit (s)'),
        ]:
            sub = paired_group_means(valid, tasks, metric, prefix)
            boxplot_groups(sub, [prefix], output, f'02_{prefix}_{suffix}',
                           ylabel, metric=metric, title=f'{prefix}: Minimum und Maximum')
    for suffix, metric, ylabel in [
        ('Prozent', 'Fehler_%', 'Durchschnittliche normalisierte Abweichung (%)'),
        ('Zeit', 'Zeit_s', 'Durchschnittliche Bearbeitungszeit (s)'),
    ]:
        sub = paired_group_means(valid, TASKS, metric, 'Gesamt')
        boxplot_groups(sub, ['Gesamt'], output, f'03_Gesamt_{suffix}',
                       ylabel, metric=metric, title='Gesamtvergleich')
    print(f'Nur completed: {len(keep)} Personen; gültige Antworten {len(valid)}/{len(trials)}; vollständig (12) {len(overall)}')
    print('Ausgaben:', output)
if __name__ == '__main__':
    main()
