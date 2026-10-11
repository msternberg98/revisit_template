"""Feedbackauswertung: Rangfolge, sechs Likert-Antworten und allgemeiner Freitext.

Ausgabe neben dem Skript unter Analysis_Output/Feedback.

Benötigt: pandas, matplotlib, openpyxl; optional scipy für statistische Tests.

"""

from pathlib import Path

import json

import numpy as np

import pandas as pd

import matplotlib.pyplot as plt

from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import re



# 1. Konfiguration

CSV_FILENAME = "Nutzerstudie_all_tidy.csv"

# CSV_FILENAME = "Nutzerstudie_all_tidy_Test.csv"

INCLUDE_INCOMPLETE = False

OUTPUT_FOLDER = Path ("Analysis_Output") / "Feedback"

SHOW_PLOTS = False

METHODS = {"Darstellung A": "VSUP", "Darstellung B": "Scaled Glyph", "Darstellung C": "IsoGlyph"}

# Gleiche Platzfarben in beiden Diagrammen (auch bei umgekehrter Stapelreihenfolge).

PLACE_COLORS = {1: "#3676b6", 2: "#e5a343", 3: "#779b7b"}



LIKERT_PROMPTS = {

    "Ablesbarkeit": {

        "VSUP": "Wie leicht konnten Sie die Werte aus der Darstellung A ablesen?",

        "Scaled Glyph": "Wie leicht konnten Sie die Werte aus der Darstellung B ablesen?",

        "IsoGlyph": "Wie leicht konnten Sie die Werte aus der Darstellung C ablesen?",

    },

    "Anstrengung": {

        "VSUP": "Wie anstrengend war die Arbeit mit der Darstellung A?",

        "Scaled Glyph": "Wie anstrengend war die Arbeit mit der Darstellung B?",

        "IsoGlyph": "Wie anstrengend war die Arbeit mit der Darstellung C?",

    },

}





def save_figure (fig, folder, stem):

    """Speichert eine Grafik als hochauflösendes PNG und als Vektor-PDF."""

    # Rechts festen Platz für die außerhalb der Achse stehende Legende lassen.

    fig.tight_layout (rect = (0, 0, 0.81, 1))

    fig.savefig (folder / f"{stem}.png", dpi=300, bbox_inches = "tight")

    fig.savefig (folder / f"{stem}.pdf", bbox_inches = "tight")

    if SHOW_PLOTS:

        plt.show ()

    plt.close (fig)





def parse_ranking (value):

    """ReVISit speichert Rangpositionen nullbasiert als JSON: 0 = Platz 1."""

    if pd.isna (value):

        return {}, "Keine Rangfolge"

    try:

        obj = json.loads (str (value))

    except (ValueError, TypeError):

        return {}, "Ungültiges JSON"

    if not isinstance (obj, dict):

        return {}, "Rangfolge ist kein JSON-Objekt"

    result = {}

    for key, rank in obj.items ():

        if key not in METHODS:

            return {}, f"Unbekannte Methode: {key}"

        try:

            position = int(rank) + 1

        except (TypeError, ValueError):

            return {}, f"Ungültiger Rang für {key}: {rank}"

        if str(rank) != str(position - 1) or position not in (1, 2, 3):

            return {}, f"Ungültiger Rang für {key}: {rank}"

        result [METHODS [key]] = position

    if len (set (result.values ())) != len (result):

        return {}, "Mehrfach vergebene Rangposition"

    return result, "Vollständig" if len (result) == 3 else "Unvollständig"





IDENT_TASKS = [
    "Niederschlag_Uncertainty_Maxima", "Niederschlag_Uncertainty_Minima",
    "Luftdruck_Value_Maxima", "Luftdruck_Value_Minima",
]
IDENT_PATTERN = re.compile(
    r"^(Niederschlag_Uncertainty_(?:Maxima|Minima)|Luftdruck_Value_(?:Maxima|Minima))_(VSUP|ScaledGlyph|IsoGlyph)$"
)
IDENT_METHODS = ["VSUP", "ScaledGlyph", "IsoGlyph"]


def personal_identification_ranks(df):
    """Dieselbe Fehlerdefinition und Complete-Case-Regel wie die Identifikationsauswertung."""
    rows = []
    for r in df.to_dict("records"):
        match = IDENT_PATTERN.fullmatch(str(r.get("trialId", "")))
        if not match:
            continue
        task, method = match.groups()
        correct = pd.to_numeric(r.get("correctAnswer"), errors="coerce")
        if pd.isna(correct):
            continue
        raw = pd.to_numeric(r.get("answer"), errors="coerce")
        valid = pd.notna(raw) and np.isfinite(raw) and raw != -999
        rows.append({"participantId": str(r["participantId"]), "Aufgabe": task,
                     "Methode": method, "Sollwert": float(correct),
                     "Antwort": float(raw) if valid else np.nan})
    cols = ["participantId"] + [f"Objektiv_Rang_{m}" for m in IDENT_METHODS] + [f"Fehler_Mittel_{m}_Prozent" for m in IDENT_METHODS]
    if not rows:
        return pd.DataFrame(columns=cols)
    trials = pd.DataFrame(rows)
    if trials.duplicated(["participantId", "Aufgabe", "Methode"]).any():
        raise ValueError("Doppelte Identifikationsantworten bei der Feedback-Verknüpfung")
    trials["Variable"] = trials["Aufgabe"].str.replace(r"_(?:Maxima|Minima)$", "", regex=True)
    references = trials.groupby(["Variable", "Aufgabe"])["Sollwert"].agg(["min", "max"])
    if (references["max"] - references["min"] > 1e-9).any():
        raise ValueError("Widersprüchliche Sollwerte bei Identifikationsaufgaben")
    spans = trials.groupby("Variable")["Sollwert"].agg(["min", "max"])
    spans["Spanne"] = spans["max"] - spans["min"]
    if (spans["Spanne"] <= 0).any():
        raise ValueError("Identifikations-Wertspanne muss positiv sein")
    trials["Fehler_%"] = 100 * (trials["Antwort"] - trials["Sollwert"]).abs() / trials["Variable"].map(spans["Spanne"])
    wide = trials.pivot(index=["participantId", "Aufgabe"], columns="Methode", values="Fehler_%").reindex(columns=IDENT_METHODS).dropna()
    counts = wide.reset_index().groupby("participantId")["Aufgabe"].nunique()
    full_ids = counts[counts.eq(len(IDENT_TASKS))].index
    wide = wide.loc[wide.index.get_level_values("participantId").isin(full_ids)]
    if wide.empty:
        return pd.DataFrame(columns=cols)
    means = wide.groupby(level="participantId").mean()
    ranks = means.rank(axis=1, method="average", ascending=True)
    out = pd.DataFrame({"participantId": means.index})
    for m in IDENT_METHODS:
        out[f"Objektiv_Rang_{m}"] = ranks[m].to_numpy()
        out[f"Fehler_Mittel_{m}_Prozent"] = means[m].to_numpy()
    return out


def compare_subjective_objective(rank_df, df):
    """Subjektive Ränge gegen objektive mittlere Identifikationsfehler."""
    objective = personal_identification_ranks(df)
    rank_df = rank_df.rename(columns={"Scaled Glyph": "ScaledGlyph"})
    result = rank_df.merge(objective, on="participantId", how="left", validate="one_to_one")
    for m in IDENT_METHODS:
        result.rename(columns={m: f"Subjektiv_Rang_{m}"}, inplace=True)
        a = result[f"Subjektiv_Rang_{m}"]
        b = result[f"Objektiv_Rang_{m}"]
        result[f"Rangdifferenz_Sub_minus_Obj_{m}"] = a - b
        result[f"Rang_uebereinstimmend_{m}"] = (a == b).where(a.notna() & b.notna(), pd.NA)
    result["Vergleich_Status"] = np.where(result["Status"].ne("Vollständig"),
                                            "Subjektive Rangfolge unvollständig/ungültig",
                                            np.where(result[f"Objektiv_Rang_{IDENT_METHODS[0]}"].isna(),
                                                     "Keine 12 gültigen Identifikationsantworten", "Vergleichbar"))
    result["Anzahl_exakt_uebereinstimmend"] = pd.NA
    result["Subjektiver_Erstplatz_objektiv_am_besten"] = pd.NA
    for idx in result.index[result["Vergleich_Status"].eq("Vergleichbar")]:
        row = result.loc[idx]
        result.at[idx, "Anzahl_exakt_uebereinstimmend"] = sum(bool(row[f"Rang_uebereinstimmend_{m}"]) for m in IDENT_METHODS)
        chosen = next(m for m in IDENT_METHODS if row[f"Subjektiv_Rang_{m}"] == 1)
        # Bei objektiven Gleichständen gilt jede Methode mit minimalem Fehler als beste.
        min_error = min(row[f"Fehler_Mittel_{m}_Prozent"] for m in IDENT_METHODS)
        result.at[idx, "Subjektiver_Erstplatz_objektiv_am_besten"] = bool(np.isclose(row[f"Fehler_Mittel_{chosen}_Prozent"], min_error, atol=1e-10, rtol=0))
    return result


def write_feedback_sheet(writer, sheetname, sections, note):
    """Lesbare Abschnittsblöcke im Stil der Identifikationsauswertung."""
    ws = writer.book.create_sheet(sheetname)
    ws.sheet_view.showGridLines = False
    ncols = max([2] + [len(frame.columns) for _, frame, _ in sections if frame is not None])
    widths = [18.] * ncols
    for _, frame, _ in sections:
        if frame is None:
            continue
        for j, col in enumerate(frame.columns):
            samples = [str(col)] + [str(v) for v in frame.iloc[:150, j] if pd.notna(v)]
            widths[j] = max(widths[j], min(55., max(map(len, samples), default=15) + 3.))
    if sheetname == "00_Lesehilfe":
        widths[0], widths[1] = 33, 105
    for j, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    def banner(row, value, color, fg="FFFFFF"):
        for j in range(1, ncols+1):
            ws.cell(row, j).fill = PatternFill("solid", fgColor=color)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell = ws.cell(row, 1, value)
        cell.font = Font(bold=(fg == "FFFFFF"), color=fg, size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 40 if len(value) > 125 else 30
    banner(1, note, "23415C")
    row = 3
    for title, frame, desc in sections:
        banner(row, title, "3676B6")
        row += 1
        if desc:
            banner(row, desc, "EAF0F6", "20354A")
            row += 1
        if frame is not None and len(frame.columns):
            for j, name in enumerate(frame.columns, 1):
                cell = ws.cell(row, j, name)
                cell.font = Font(bold=True, color="20354A")
                cell.fill = PatternFill("solid", fgColor="D7E5F1")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[row].height = 42
            row += 1
            for values in frame.itertuples(index=False, name=None):
                for j, val in enumerate(values, 1):
                    if pd.isna(val):
                        val = None
                    if isinstance(val, np.generic):
                        val = val.item()
                    if isinstance(val, (bool, np.bool_)):
                        val = "Ja" if val else "Nein"
                    cell = ws.cell(row, j, val)
                    is_number = isinstance(val, (int, float)) and not isinstance(val, bool)
                    cell.alignment = Alignment(horizontal="center" if is_number else "left", vertical="center", wrap_text=True)
                    if is_number:
                        # Rangpositionen und Häufigkeiten als ganze Zahlen; Mittelränge dürfen .5 enthalten.
                        if str(frame.columns[j - 1]).startswith("Subjektiv_Rang_"):
                            cell.number_format = '0'
                        elif isinstance(val, int) or (isinstance(val, float) and val.is_integer() and str(frame.columns[j - 1]).startswith(("Platz ", "Anzahl", "n"))):
                            cell.number_format = '0'
                        else:
                            cell.number_format = '#,##0.0000;[Red]-#,##0.0000' 
                ws.row_dimensions[row].height = 28
                row += 1
        row += 2


def rank_plots (counts, folder):

    """Gruppierte und gestapelte Säulen mit identischen Platzfarben."""

    names = list (METHODS.values())

    x = np.arange (len (names))

    fig, ax = plt.subplots (figsize = (8, 5))

    for place, offset in [(1, -0.25), (2, 0), (3, 0.25)]:

        values = [int (counts.loc [name, place]) for name in names]

        bars = ax.bar (x + offset, values, width = 0.24, color = PLACE_COLORS [place], label = f"Platz {place}")

        ax.bar_label (bars, padding = 3, fmt = "%d")

    ax.set_xticks (x, names)

    ax.set_ylabel ("Anzahl der Platzierungen")

    ax.set_title ("Platzierungsverteilung (gruppiert)")

    ax.yaxis.get_major_locator ().set_params (integer = True)

    ax.set_ylim (0, max (1, int (counts.to_numpy ().max ())) * 1.25 + 0.5)

    ax.legend (title = "Platzierung", loc = "center left", bbox_to_anchor = (1.02, 0.5), borderaxespad = 0)

    ax.spines [["top", "right"]].set_visible (False)

    save_figure (fig, folder, "01_Rangfolge_Gruppiert")



    fig, ax = plt.subplots (figsize = (8, 5))

    bottoms = np.zeros (len (names), dtype = int)

    # Von unten nach oben: Platz 3, Platz 2, Platz 1.

    for place in (3, 2, 1):

        values = np.array ([int (counts.loc [name, place]) for name in names])

        bars = ax.bar (names, values, bottom = bottoms, color = PLACE_COLORS [place], label = f"Platz {place}")

        ax.bar_label (bars, label_type = "center", labels = [str (v) if v else "" for v in values])

        bottoms += values

    ax.set_ylabel ("Anzahl der Platzierungen")

    ax.set_title ("Platzierungsverteilung (gestapelt)")

    ax.yaxis.get_major_locator ().set_params (integer = True)

    ax.set_ylim (0, max (1, int (bottoms.max ())) * 1.15 + 0.5)

    handles, labels = ax.get_legend_handles_labels ()

    ax.legend (handles [::-1], labels [::-1], title = "Platzierung", loc = "center left", bbox_to_anchor = (1.02, 0.5), borderaxespad = 0)

    ax.spines [["top", "right"]].set_visible (False)

    save_figure (fig, folder, "02_Rangfolge_Gestapelt")





def likert_plot (counts, dimension, folder):

    """Zeigt alle sieben Stufen je Methode als gruppierte Säulen."""

    names = list (METHODS.values ())

    fig, ax = plt.subplots (figsize = (10, 5))

    x = np.arange (1, 8)

    for i, method in enumerate (names):

        values = counts.loc [method, list (range(1, 8))].to_numpy (dtype = int)

        ax.bar (x + (i - 1) * 0.24, values, width = 0.23, color = PLACE_COLORS [i + 1], label = method)

    ax.set_xticks (x, [str (v) for v in x])

    ax.set_xlabel ("Antwortstufe (1-7)")

    ax.set_ylabel ("Anzahl der Antworten")

    ax.set_title (dimension)

    ax.yaxis.get_major_locator ().set_params (integer=True)

    ax.legend (title = "Methode", loc = "center left", bbox_to_anchor = (1.02, 0.5), borderaxespad = 0)

    ax.spines [["top", "right"]].set_visible (False)

    save_figure (fig, folder, "03_Ablesbarkeit" if dimension == "Ablesbarkeit" else "04_Anstrengung")





def run_tests (wide, label):

    """Friedman-Test; bei p<0,05 paarweise Wilcoxon-Tests mit Holm-Korrektur."""

    methods = list (METHODS.values ())

    paired = wide.reindex (columns = methods).dropna ()

    results = [{"Merkmal": label, "Test": "Vollständige Datensätze", "n": len (paired)}]

    if len (paired) < 3:

        results.append ({"Merkmal": label, "Test": "Hinweis", "Ergebnis": "Weniger als drei vollständige Datensätze"})

        return results

    try:

        from scipy.stats import friedmanchisquare, wilcoxon

    except ImportError:

        results.append ({"Merkmal": label, "Test": "Hinweis", "Ergebnis": "scipy nicht installiert"})

        return results

    if all (paired [m].equals (paired [methods [0]]) for m in methods [1:]):

        results.append ({"Merkmal": label, "Test": "Hinweis", "Ergebnis": "Alle Methoden identisch bewertet; kein Test"})

        return results

    statistic, p = friedmanchisquare (*(paired [m] for m in methods))

    results.append ({"Merkmal": label, "Test": "Friedman", "n": len (paired), "Statistik": statistic, "p": p, "Signifikanz (p < 0.05)": "Signifikant" if p < 0.05 else "Nicht signifikant",})

    if not np.isfinite (p) or p >= 0.05:

        return results

    pairs = [(methods [0], methods [1]), (methods [0], methods [2]), (methods [1], methods [2])]

    raw = []

    for a, b in pairs:

        differences = paired [a] - paired [b]

        if (differences == 0).all ():

            stat, pv = 0.0, 1.0

        else:

            stat, pv = wilcoxon (paired [a], paired [b], zero_method = "wilcox", alternative = "two-sided")

        raw.append ((a, b, stat, pv))

    # Holm: sortierte p-Werte, monoton korrigierte Werte.

    ordering = sorted (range (3), key = lambda i: raw [i] [3])

    adjusted = [None] * 3

    running = 0.0

    for j, idx in enumerate (ordering):

        running = max (running, min (1.0, raw [idx] [3] * (3 - j)))

        adjusted [idx] = running

    for i, (a, b, stat, pv) in enumerate (raw):

        results.append ({"Merkmal": label, "Test": "Wilcoxon (Holm)", "Vergleich": f"{a} - {b}",

                        "n": len (paired), "Statistik": stat, "p": pv, "p_Holm": adjusted [i]})

    return results





def main ():

    # 2. CSV laden und wie bei der Demografie nur abgeschlossene Teilnahmen auswählen.

    folder = Path (__file__).resolve ().parent

    source = folder / CSV_FILENAME

    if not source.exists ():

        raise FileNotFoundError (f"CSV nicht gefunden: {source}")

    out = folder / OUTPUT_FOLDER

    out.mkdir (parents = True, exist_ok = True)

    df = pd.read_csv (source, low_memory = False, dtype = {"participantId": str})

    required = {"participantId", "percentComplete", "trialId", "status", "responsePrompt", "answer"}

    if missing := required.difference (df.columns):

        raise ValueError (f"Fehlende CSV-Spalten: {sorted (missing)}")

    if not INCLUDE_INCOMPLETE:

        pc = pd.to_numeric (df ["percentComplete"], errors = "coerce")



        participant_info = df.assign (_pc = pc).groupby ("participantId").agg (

            max_pc = ("_pc", "max"),

            has_completed = ("status", lambda s: s.eq ("completed").any ()),

            has_rejected = ("status", lambda s: s.eq ("rejected").any ()),

        )



        keep = participant_info.index [

            ~participant_info ["has_rejected"]

            & participant_info ["has_completed"]

        ]



        df = df [df ["participantId"].isin (keep)].copy ()

    ids = sorted (df ["participantId"].dropna ().unique ())

    feedback = df [df ["trialId"] == "Feedback"].copy ()



    # 3. Rangfolge: ReVISit speichert 0/1/2, nicht 1/2/3.

    ranking_rows = feedback [feedback ["responsePrompt"].fillna ("").str.startswith ("Welche Darstellung fanden Sie am hilfreichsten")]

    if ranking_rows.duplicated ("participantId").any ():

        raise ValueError ("Mehrere Rangfolgen pro Person gefunden; Rohdaten prüfen.")

    rankings = []

    for pid in ids:

        selected = ranking_rows.loc [ranking_rows ["participantId"] == pid, "answer"]

        parsed, status = parse_ranking (selected.iloc [0]) if not selected.empty else ({}, "Keine Rangfolge")

        rankings.append ({"participantId": pid, ** {m: parsed.get (m, np.nan) for m in METHODS.values ()}, "Status": status})

    rank_df = pd.DataFrame (rankings)

    complete = rank_df [rank_df ["Status"] == "Vollständig"]

    if (rank_df ["Status"] != "Vollständig").any ():

        print ("WARNUNG: Unvollständige/ungültige Rangfolgen; für Rangtests ausgeschlossen:")

        print (rank_df [rank_df ["Status"] != "Vollständig"] [['participantId', 'Status']].to_string (index=False))

    counts = pd.DataFrame (index = list (METHODS.values ()), columns = [1, 2, 3], data = 0)

    # Häufigkeitsdiagramme basieren ausschließlich auf vollständigen Rangfolgen.

    for method in METHODS.values ():

        for place in (1, 2, 3):

            counts.loc [method, place] = int ((complete [method] == place).sum ())

    counts.index.name = "Methode"

    counts.columns = [f"Platz {i}" for i in (1, 2, 3)]

    export_counts = counts.reset_index ()

    counts.columns = [1, 2, 3]

    rank_summary = pd.DataFrame ({"Methode": list (METHODS.values ()),

                                 "Mittlerer Rang": [complete [m].mean () for m in METHODS.values ()],

                                 "Vollständige Rangfolgen": len (complete),

                                 "Unvollständige Rangfolgen": len (rank_df) - len (complete)})

    rank_combined = export_counts.merge (rank_summary, on = "Methode", validate = "one_to_one")

    rank_combined.to_csv (out / "01_Rangfolge_Haeufigkeiten.csv", sep = ";", decimal = ",", encoding = "utf-8-sig", index = False)

    rank_compare = compare_subjective_objective(rank_df, df)
    rank_compare.to_csv (out / "01_Rangfolge_Einzelantworten.csv", sep = ";", decimal = ",", encoding = "utf-8-sig", index = False)

    rank_plots (counts, out)

    tests = run_tests (complete.set_index ("participantId"), "Rangfolge")



    # 4. Likert-Fragen: jede der beiden Fragen getrennt, alle Stufen 1–7.

    likert_tables = {}

    for dimension, prompts in LIKERT_PROMPTS.items ():

        rows = []

        stats_rows = []

        wide = pd.DataFrame (index=ids)

        for method, prompt in prompts.items ():

            matches = feedback [feedback ["responsePrompt"] == prompt]

            if matches.duplicated ("participantId").any ():

                raise ValueError (f"Mehrfachantworten für {method}, {dimension}")

            values = pd.to_numeric (matches.set_index ("participantId") ["answer"], errors = "coerce").reindex (ids)

            invalid = values.notna () & ~values.isin (range(1, 8))

            if invalid.any ():

                raise ValueError (f"Ungültige Likert-Stufen in {dimension} / {method}")

            wide [method] = values

            valid = values.dropna ()

            for level in range (1, 8):

                n = int ((valid == level).sum ())

                rows.append ({"Methode": method, "Stufe": level, "Anzahl": n,

                             "Anteil gültiger Antworten (%)": round (100 * n / len (valid), 2) if len (valid) else 0})

            stats_rows.append ({"Methode": method, "n": len (valid),

                                 "Fehlend": len (ids) - len (valid), "Median": valid.median (),

                                 "Q1": valid.quantile (.25), "Q3": valid.quantile (.75),

                                 "IQR": valid.quantile (.75) - valid.quantile (.25),

                                 "Mittelwert": valid.mean (), "Standardabweichung": valid.std ()})

        table = pd.DataFrame (rows)

        # Eine Zeile pro Methode: Stufenhäufigkeiten, Prozentwerte und Kennzahlen.

        frequencies = table.pivot (index = "Methode", columns = "Stufe", values = "Anzahl")

        frequencies = frequencies.reindex (index = list (METHODS.values ()), columns = range (1, 8), fill_value = 0)

        percentages = table.pivot (index = "Methode", columns = "Stufe", values = "Anteil gültiger Antworten (%)")

        percentages = percentages.reindex (index = list (METHODS.values ()), columns = range (1, 8), fill_value = 0)

        # Je Methode zwei Zeilen: absolute Häufigkeiten unmittelbar über Prozentwerten.

        # Deskriptive Kennzahlen stehen ausschließlich in der Prozentzeile.

        stats_by_method = pd.DataFrame (stats_rows).set_index ("Methode")

        combined_rows = []

        for method in METHODS.values ():

            count_row = {"Methode": method, "Angabe": "Anzahl"}

            percent_row = {"Methode": method, "Angabe": "Prozent (%)"}

            for level in range (1, 8):

                count_row [f"Stufe {level}"] = int (frequencies.loc [method, level])

                percent_row [f"Stufe {level}"] = float (percentages.loc [method, level])

            for col in stats_by_method.columns:

                count_row [col] = None

                percent_row [col] = stats_by_method.loc [method, col]

            combined_rows.extend ([count_row, percent_row])

        combined = pd.DataFrame (combined_rows)

        likert_tables [dimension] = combined

        combined.to_csv (out / f"0{'2' if dimension == 'Ablesbarkeit' else '3'}_{dimension}.csv",

                     sep = ";", decimal = ",", encoding = "utf-8-sig", index = False)

        wide.index.name = "participantId"

        wide.reset_index ().to_csv (out / f"{dimension}_Einzelantworten.csv", sep = ";", decimal = ",", encoding = "utf-8-sig", index = False)

        likert_plot (table.pivot (index = "Methode", columns = "Stufe", values = "Anzahl"), dimension, out)

        tests.extend (run_tests (wide, dimension))

        if wide.isna ().any ().any ():

            print (f"WARNUNG: Fehlende Antworten bei {dimension}: {wide.isna ().sum ().to_dict ()}")



    # 5. Allgemeiner Freitext: keine automatische inhaltliche Interpretation.

    notes = feedback [feedback ["responsePrompt"] == "Weitere Anmerkungen"]

    if notes.duplicated ("participantId").any ():

        raise ValueError ("Mehrere allgemeine Freitextantworten pro Person gefunden")

    notes = notes [["participantId", "answer"]].rename (columns = {"answer": "Anmerkung"}).copy ()

    notes ["Anmerkung"] = notes ["Anmerkung"].fillna ("").astype (str).str.strip ()

    notes = notes [notes ["Anmerkung"] != ""]

    notes.to_csv (out / "04_Freitextantworten.csv", sep = ";", encoding = "utf-8-sig", index = False)



    # 6. Excel-Gesamtdatei: Abschnittsstruktur wie bei Identifikation.
    with pd.ExcelWriter (out / "00_Feedback_Auswertung.xlsx", engine = "openpyxl") as writer:
        guide = pd.DataFrame([
            ("Datenbasis", "Nur Teilnahmen mit status=completed und ohne status=rejected; laufende Teilnahmen ausgeschlossen."),
            ("Subjektiver Rang", "ReVISit-Rangfolge aus der Feedbackfrage: 1 = am hilfreichsten, 3 = am wenigsten hilfreich. Keine Gleichstände."),
            ("Mittlerer subjektiver Rang", "Arithmetisches Mittel der vergebenen Rangplätze über vollständige Feedback-Rangfolgen; kleiner = bevorzugt."),
            ("Objektiver Rang", "Pro Person und Methode: Mittelwert der normalisierten absoluten Fehler (%) über alle vier Identifikationsaufgaben. Niedrigster Mittelwert = Rang 1; Gleichstände erhalten Mittelränge."),
            ("Vollständigkeit", "Objektive Ränge nur bei 12 gültigen Antworten (4 Aufgaben × 3 Methoden). Subjektive Vergleiche nur bei vollständiger Feedback-Rangfolge."),
            ("Rangdifferenz", "Subjektiver Rang minus objektiver Rang. Null = identisch; positive Werte = subjektiv schlechter platziert als anhand des Fehlers; negative = subjektiv besser platziert."),
            ("Exakte Übereinstimmung", "Zählt die Methoden, deren subjektiver Rang exakt dem objektiven Rang entspricht (0 bis 3). Objektive Mittelränge bei Gleichstand können eine exakte Übereinstimmung verhindern."),
            ("Erstplatz-Übereinstimmung", "Ja, wenn die subjektiv auf Platz 1 gesetzte Methode den kleinsten persönlichen Durchschnittsfehler hat; objektive Gleichstände auf Platz 1 zählen ebenfalls."),
            ("Ablesbarkeit und Anstrengung", "Je Methode 7-stufige Antworten: Anzahl, Anteil gültiger Antworten und deskriptive Kennzahlen. Keine Änderung der bisherigen Berechnungen."),
            ("Friedman und Wilcoxon", "Friedman prüft Unterschiede zwischen den drei verbundenen Feedbackbewertungen; bei p < 0,05 folgen paarweise Wilcoxon-Tests mit Holm-Korrektur innerhalb des Merkmals."),
            ("Interpretation", "Die Übereinstimmung beschreibt, ob persönliche Präferenz und Identifikationsgenauigkeit zusammenfallen; sie beweist keine Ursache und ist kein zusätzlicher Signifikanztest."),
        ], columns=["Begriff", "Bedeutung / Berechnung"])
        write_feedback_sheet(writer, "00_Lesehilfe", [("Berechnungsgrundlagen", guide, "")], "Feedback-Auswertung: Erläuterungen zu den Tabellen")
        write_feedback_sheet(writer, "Rangfolge", [
            ("A. Platzierungshäufigkeiten und mittlere subjektive Ränge", rank_combined,
             "Platz 1 = am hilfreichsten; mittlerer Rang = Durchschnitt der subjektiven Plätze über vollständige Feedback-Rangfolgen.")
        ], "Subjektive Rangfolge der drei Visualisierungsmethoden")
        subjective_cols = ["participantId", "Status"] + [f"Subjektiv_Rang_{m}" for m in IDENT_METHODS]
        objective_cols = ["participantId", "Vergleich_Status"] + [f"Fehler_Mittel_{m}_Prozent" for m in IDENT_METHODS] + [f"Objektiv_Rang_{m}" for m in IDENT_METHODS]
        comparison_cols = ["participantId", "Vergleich_Status"] + [f"Rangdifferenz_Sub_minus_Obj_{m}" for m in IDENT_METHODS] + [f"Rang_uebereinstimmend_{m}" for m in IDENT_METHODS] + ["Anzahl_exakt_uebereinstimmend", "Subjektiver_Erstplatz_objektiv_am_besten"]
        write_feedback_sheet(writer, "Rang-Einzelantworten", [
            ("A. Subjektive Rangfolge", rank_compare.reindex(columns=subjective_cols),
             "Rang 1 = subjektiv am hilfreichsten. Nur vollständige Rangfolgen werden für Vergleiche verwendet."),
            ("B. Objektive Ränge aus Identifikationsaufgaben", rank_compare.reindex(columns=objective_cols),
             "Je Person: Mittel aus vier normalisierten Fehlern pro Methode; Rang 1 = kleinster Fehler. Nur vollständige 12 Identifikationsantworten. Gleichstände: Mittelränge."),
            ("C. Subjektiv-objektiver Abgleich", rank_compare.reindex(columns=comparison_cols),
             "Rangdifferenz = subjektiver Rang minus objektiver Rang. Exakte Rangübereinstimmung pro Methode; Erstplatz stimmt überein, wenn die subjektive Lieblingsmethode objektiv den geringsten Durchschnittsfehler erzielt (einschließlich Gleichstand).")
        ], "Persönlicher Vergleich von Präferenz und Identifikationsgenauigkeit")
        for name, table in likert_tables.items():
            write_feedback_sheet(writer, name, [("Antwortverteilung und Kennzahlen", table,
                "Stufe 1–7: Anzahl und prozentualer Anteil; Median, Quartile, IQR, Mittelwert und Standardabweichung basieren auf gültigen Antworten.")], f"Feedback: {name}")
        write_feedback_sheet(writer, "Statistische Tests", [("Friedman und paarweise Wilcoxon-Tests", pd.DataFrame(tests),
            "Verbundenes Design: dieselben Personen bewerten alle drei Methoden. Friedman als Gesamttest; bei p < 0,05 Wilcoxon-Paarvergleiche mit Holm-Korrektur.")], "Statistische Tests der subjektiven Feedbackbewertungen")
        write_feedback_sheet(writer, "Freitext", [("Allgemeine Anmerkungen", notes,
            "Originalantworten ohne automatische inhaltliche Auswertung.")], "Freitextantworten")
    print (f"\nTeilnehmende: {len (ids)}; vollständige Rangfolgen: {len (complete)}; Freitextantworten: {len (notes)}")

    print ("\nPlatzierungsverteilung:\n", export_counts.to_string (index = False))

    print (f"\nFertig. Dateien gespeichert unter: {out}")





if __name__ == "__main__":

    main()
